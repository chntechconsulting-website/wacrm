import io
import csv
from datetime import datetime, date
from typing import Optional, Generator, List
from fastapi import APIRouter, Depends, Query, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, and_, desc

from database import get_db
from models.user import User
from models.employee import Employee
from models.lead import Lead
from models.call_log import CallLog
from services.auth import require_roles, get_current_user

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/reports/export", tags=["Admin Data Export"])

def _format_dt(dt: Optional[datetime]) -> str:
    if not dt:
        return ""
    return dt.strftime("%Y-%m-%d %H:%M:%S")

def _format_date(d: Optional[date]) -> str:
    if not d:
        return ""
    return d.strftime("%Y-%m-%d")

# -----------------------------------------------------------------------------
# 1. Export Count / Stats Endpoint
# -----------------------------------------------------------------------------
@router.get("/stats")
def get_export_stats(
    lead_status: Optional[str] = None,
    lead_start_date: Optional[str] = None,
    lead_end_date: Optional[str] = None,
    call_status: Optional[str] = None,
    call_employee_id: Optional[int] = None,
    call_start_date: Optional[str] = None,
    call_end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["ADMIN"]))
):
    """
    Returns the exact number of matching leads and call logs based on applied filters.
    """
    # Leads count
    leads_q = db.query(Lead)
    if lead_status and lead_status.strip() and lead_status.upper() != "ALL":
        leads_q = leads_q.filter(Lead.status.ilike(f"%{lead_status.strip()}%"))
    if lead_start_date:
        try:
            s_dt = datetime.strptime(lead_start_date, "%Y-%m-%d")
            leads_q = leads_q.filter(Lead.created_at >= s_dt)
        except ValueError:
            pass
    if lead_end_date:
        try:
            e_dt = datetime.strptime(lead_end_date + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            leads_q = leads_q.filter(Lead.created_at <= e_dt)
        except ValueError:
            pass
    total_leads = leads_q.count()

    # Call logs count
    calls_q = db.query(CallLog)
    if call_status and call_status.strip() and call_status.upper() != "ALL":
        calls_q = calls_q.filter(CallLog.call_status.ilike(f"%{call_status.strip()}%"))
    if call_employee_id and call_employee_id > 0:
        calls_q = calls_q.filter(CallLog.employee_id == call_employee_id)
    if call_start_date:
        try:
            s_dt = datetime.strptime(call_start_date, "%Y-%m-%d")
            calls_q = calls_q.filter(CallLog.created_at >= s_dt)
        except ValueError:
            pass
    if call_end_date:
        try:
            e_dt = datetime.strptime(call_end_date + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            calls_q = calls_q.filter(CallLog.created_at <= e_dt)
        except ValueError:
            pass
    total_calls = calls_q.count()

    return {
        "matching_leads": total_leads,
        "matching_calls": total_calls,
        "total_database_leads": db.query(Lead).count(),
        "total_database_calls": db.query(CallLog).count()
    }


# -----------------------------------------------------------------------------
# 2. ALL LEADS EXPORT
# -----------------------------------------------------------------------------
LEAD_COLUMNS = [
    ("ID", "id"),
    ("Name", "name"),
    ("First Name", "first_name"),
    ("Last Name", "last_name"),
    ("Phone", "phone"),
    ("WhatsApp Number", "whatsapp_number"),
    ("Email", "email"),
    ("Designation", "designation"),
    ("Company", "company"),
    ("Industry", "industry"),
    ("Website", "website"),
    ("Company Size", "company_size"),
    ("City", "city"),
    ("State", "state"),
    ("Country", "country"),
    ("Location", "location"),
    ("Source", "source"),
    ("Campaign Name", "campaign_name"),
    ("Meta Lead ID", "meta_lead_id"),
    ("Meta Form ID", "meta_form_id"),
    ("Website Name", "website_name"),
    ("Landing Page", "landing_page"),
    ("Status", "status"),
    ("Rating", "rating"),
    ("Priority", "priority"),
    ("Assigned Employee ID", "assigned_to"),
    ("Assigned Employee Name", "assigned_emp_name"),
    ("Assigned At", "assigned_at"),
    ("Requirement Notes", "requirement_notes"),
    ("Follow-up Date", "followup_date"),
    ("Follow-up Time", "followup_time"),
    ("QuoteTask Quotation Number", "quotetask_quotation_number"),
    ("QuoteTask Quotation Status", "quotetask_quotation_status"),
    ("QuoteTask Quotation Amount", "quotetask_quotation_amount"),
    ("Is Re-Inquiry", "is_reinquiry"),
    ("Re-Inquiry Count", "reinquiry_count"),
    ("Total Enquiries", "total_enquiries"),
    ("First Enquiry Date", "first_enquiry_date"),
    ("Latest Enquiry Date", "latest_enquiry_date"),
    ("Created At", "created_at"),
    ("Updated At", "updated_at")
]

def _extract_lead_row(lead: Lead, emp_map: dict) -> list:
    emp_name = emp_map.get(lead.assigned_to, "") if lead.assigned_to else ""
    return [
        lead.id,
        lead.name or "",
        lead.first_name or "",
        lead.last_name or "",
        lead.phone or "",
        lead.whatsapp_number or "",
        lead.email or "",
        lead.designation or "",
        lead.company or "",
        lead.industry or "",
        lead.website or "",
        lead.company_size or "",
        lead.city or "",
        lead.state or "",
        lead.country or "India",
        lead.location or "",
        lead.source or "",
        lead.campaign_name or "",
        lead.meta_lead_id or "",
        lead.meta_form_id or "",
        lead.website_name or "",
        lead.landing_page or "",
        lead.status or "",
        lead.rating or "",
        lead.priority or "",
        lead.assigned_to or "",
        emp_name,
        _format_dt(lead.assigned_at),
        lead.requirement_notes or "",
        _format_date(lead.followup_date),
        lead.followup_time or "",
        lead.quotetask_quotation_number or "",
        lead.quotetask_quotation_status or "",
        lead.quotetask_quotation_amount if lead.quotetask_quotation_amount is not None else "",
        "YES" if lead.is_reinquiry else "NO",
        lead.reinquiry_count if lead.reinquiry_count is not None else 0,
        lead.total_enquiries if lead.total_enquiries is not None else 1,
        _format_dt(lead.first_enquiry_date),
        _format_dt(lead.latest_enquiry_date),
        _format_dt(lead.created_at),
        _format_dt(lead.updated_at)
    ]

@router.get("/leads")
def export_all_leads(
    format: str = "xlsx",
    status: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["ADMIN"]))
):
    """
    Exports ALL leads matching the criteria into Excel (.xlsx) or CSV format.
    If no filters selected, streams all available leads in the database.
    """
    query = db.query(Lead)

    if status and status.strip() and status.upper() != "ALL":
        query = query.filter(Lead.status.ilike(f"%{status.strip()}%"))

    if start_date:
        try:
            s_dt = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(Lead.created_at >= s_dt)
        except ValueError:
            pass

    if end_date:
        try:
            e_dt = datetime.strptime(end_date + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            query = query.filter(Lead.created_at <= e_dt)
        except ValueError:
            pass

    query = query.order_by(desc(Lead.id))
    all_leads = query.all()

    # Pre-fetch employees map to avoid N+1 query overhead
    employees = db.query(Employee).all()
    emp_map = {e.id: e.name for e in employees}

    today_str = datetime.utcnow().strftime("%Y-%m-%d")

    # CSV Generation
    if format == "csv":
        def csv_generator() -> Generator[str, None, None]:
            output = io.StringIO()
            writer = csv.writer(output)
            # Write header
            writer.writerow([col[0] for col in LEAD_COLUMNS])
            yield output.getvalue()
            output.seek(0)
            output.truncate(0)

            for lead in all_leads:
                writer.writerow(_extract_lead_row(lead, emp_map))
                yield output.getvalue()
                output.seek(0)
                output.truncate(0)

        filename = f"WACRM_All_Leads_{today_str}.csv"
        return StreamingResponse(
            csv_generator(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    # Excel Generation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Leads"
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Write headers
    headers = [col[0] for col in LEAD_COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    # Write data rows
    for row_idx, lead in enumerate(all_leads, start=2):
        row_data = _extract_lead_row(lead, emp_map)
        ws.append(row_data)
        # Apply zebra striping & borders
        fill = PatternFill(start_color="F8FAFC" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = align_left

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"WACRM_All_Leads_{today_str}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# -----------------------------------------------------------------------------
# 3. ALL CALL LOGS EXPORT
# -----------------------------------------------------------------------------
CALL_COLUMNS = [
    ("Call ID", "id"),
    ("Provider", "provider"),
    ("Provider Call Reference ID", "provider_call_id"),
    ("Employee ID", "employee_id"),
    ("Employee Name", "employee_name"),
    ("Lead ID", "lead_id"),
    ("Lead Name", "lead_name"),
    ("Lead Company", "lead_company"),
    ("From Number", "from_number"),
    ("To Number", "to_number"),
    ("Direction", "call_direction"),
    ("Call Status", "call_status"),
    ("Duration (Seconds)", "duration_seconds"),
    ("Duration (Minutes)", "duration_minutes"),
    ("Started At", "started_at"),
    ("Answered At", "answered_at"),
    ("Ended At", "ended_at"),
    ("Call Notes", "notes"),
    ("AI Sentiment", "sentiment"),
    ("AI Summary", "ai_summary"),
    ("Transcript", "transcript"),
    ("Recording URL", "recording_url"),
    ("Logged At", "created_at")
]

def _extract_call_row(call: CallLog, emp_map: dict, lead_map: dict) -> list:
    emp_name = emp_map.get(call.employee_id, "") if call.employee_id else ""
    lead_info = lead_map.get(call.lead_id, {}) if call.lead_id else {}
    duration_mins = round((call.duration_seconds or 0) / 60.0, 2)
    return [
        call.id,
        call.provider or "",
        call.provider_call_id or "",
        call.employee_id or "",
        emp_name,
        call.lead_id or "",
        lead_info.get("name", ""),
        lead_info.get("company", ""),
        call.from_number or "",
        call.to_number or "",
        call.call_direction or "OUTBOUND",
        call.call_status or "",
        call.duration_seconds or 0,
        duration_mins,
        _format_dt(call.started_at),
        _format_dt(call.answered_at),
        _format_dt(call.ended_at),
        call.notes or "",
        call.sentiment or "",
        call.ai_summary or "",
        call.transcript or "",
        call.recording_url or "",
        _format_dt(call.created_at)
    ]

@router.get("/call-logs")
def export_all_call_logs(
    format: str = "xlsx",
    status: Optional[str] = None,
    employee_id: Optional[int] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["ADMIN"]))
):
    """
    Exports ALL call logs matching the criteria into Excel (.xlsx) or CSV format.
    If no filters selected, exports all call log history in the database.
    """
    query = db.query(CallLog)

    if status and status.strip() and status.upper() != "ALL":
        query = query.filter(CallLog.call_status.ilike(f"%{status.strip()}%"))

    if employee_id and employee_id > 0:
        query = query.filter(CallLog.employee_id == employee_id)

    if start_date:
        try:
            s_dt = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(CallLog.created_at >= s_dt)
        except ValueError:
            pass

    if end_date:
        try:
            e_dt = datetime.strptime(end_date + " 23:59:59", "%Y-%m-%d %H:%M:%S")
            query = query.filter(CallLog.created_at <= e_dt)
        except ValueError:
            pass

    query = query.order_by(desc(CallLog.id))
    all_calls = query.all()

    # Pre-fetch employees and leads maps
    employees = db.query(Employee).all()
    emp_map = {e.id: e.name for e in employees}

    lead_ids = [c.lead_id for c in all_calls if c.lead_id]
    leads = db.query(Lead).filter(Lead.id.in_(lead_ids)).all() if lead_ids else []
    lead_map = {l.id: {"name": l.name, "company": l.company or ""} for l in leads}

    today_str = datetime.utcnow().strftime("%Y-%m-%d")

    # CSV Generation
    if format == "csv":
        def csv_generator() -> Generator[str, None, None]:
            output = io.StringIO()
            writer = csv.writer(output)
            # Write header
            writer.writerow([col[0] for col in CALL_COLUMNS])
            yield output.getvalue()
            output.seek(0)
            output.truncate(0)

            for call in all_calls:
                writer.writerow(_extract_call_row(call, emp_map, lead_map))
                yield output.getvalue()
                output.seek(0)
                output.truncate(0)

        filename = f"WACRM_All_Call_Logs_{today_str}.csv"
        return StreamingResponse(
            csv_generator(),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    # Excel Generation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Call Logs"
    ws.views.sheetView[0].showGridLines = True

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")  # Teal brand header
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Write headers
    headers = [col[0] for col in CALL_COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    # Write data rows
    for row_idx, call in enumerate(all_calls, start=2):
        row_data = _extract_call_row(call, emp_map, lead_map)
        ws.append(row_data)
        fill = PatternFill(start_color="F0FDFA" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = align_left

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"WACRM_All_Call_Logs_{today_str}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
